"""
Builds the admin panel guide as a Word document (.docx).

The words live in admin-panel-guide.py (sections 1-6, between the
"contents" and "build" markers). This script runs that same content
through Word-flavoured versions of its helpers (p, h1, bullets, table, box…),
so the Word file and the PDF always say the same thing.

    python docs/client-kit/admin-panel-guide-word.py docs/client-kit/admin-panel-guide.docx

Needs `pip install python-docx`. To make the PDF from it, open the .docx in
Word and Save As PDF (or see README: Word does the conversion).
"""
import html
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

OUT = sys.argv[1]
SOURCE = Path(__file__).with_name('admin-panel-guide.py').read_text(encoding='utf-8')

BURGUNDY = RGBColor(0x9B, 0x2C, 0x40)
INK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x66, 0x66, 0x66)
FONT = 'Arial'

SITE_URL = re.search(r"SITE_URL = '([^']+)'", SOURCE).group(1)
ADMIN_URL = re.search(r"ADMIN_URL = '([^']+)'", SOURCE).group(1)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
for side in ('left_margin', 'right_margin'):
    setattr(sec, side, Mm(18))
sec.top_margin, sec.bottom_margin = Mm(18), Mm(16)
CONTENT_W = Mm(210 - 36)

normal = doc.styles['Normal']
normal.font.name = FONT
normal.font.size = Pt(10)
normal.font.color.rgb = INK
normal.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.2

for name, size, color, before, after in [
    ('Heading 1', 20, BURGUNDY, 0, 10),
    ('Heading 2', 13.5, INK, 12, 5),
    ('Heading 3', 11, BURGUNDY, 8, 3),
]:
    st = doc.styles[name]
    st.font.name = FONT
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = color
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    st.element.rPr.rFonts.set(qn('w:ascii'), FONT)
    st.element.rPr.rFonts.set(qn('w:hAnsi'), FONT)
    # Word's heading styles point at the theme font, which wins over the names above.
    for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        st.element.rPr.rFonts.attrib.pop(qn(attr), None)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True


# ---------------------------------------------------------------- inline markup
TOKEN = re.compile(r'(<link href="[^"]*"[^>]*>|</link>|</?b>|</?i>|</?u>)')


def add_hyperlink(par, url, text, bold=False, italic=False):
    rel = par.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), rel)
    run = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    for tag, val in [('w:color', '9B2C40'), ('w:u', 'single')]:
        el = OxmlElement(tag)
        el.set(qn('w:val'), val)
        rpr.append(el)
    if bold:
        rpr.append(OxmlElement('w:b'))
    if italic:
        rpr.append(OxmlElement('w:i'))
    run.append(rpr)
    t = OxmlElement('w:t')
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    run.append(t)
    link.append(run)
    par._p.append(link)


def rich(par, markup, size=None, color=None):
    """Writes reportlab-style markup (<b>, <i>, <u>, <link>) into a Word paragraph."""
    bold = italic = under = False
    href = None
    for part in TOKEN.split(markup):
        if not part:
            continue
        if part == '<b>':
            bold = True
        elif part == '</b>':
            bold = False
        elif part == '<i>':
            italic = True
        elif part == '</i>':
            italic = False
        elif part == '<u>':
            under = True
        elif part == '</u>':
            under = False
        elif part.startswith('<link'):
            href = re.search(r'href="([^"]*)"', part).group(1)
        elif part == '</link>':
            href = None
        else:
            text = html.unescape(part)
            if href:
                add_hyperlink(par, href, text, bold, italic)
                continue
            run = par.add_run(text)
            run.bold, run.italic, run.underline = bold or None, italic or None, under or None
            if size:
                run.font.size = Pt(size)
            if color is not None:
                run.font.color.rgb = color
    return par


def shade(cell, hex_fill):
    tcpr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto')
    sh.set(qn('w:fill'), hex_fill)
    tcpr.append(sh)


def borders(table, color='E6E0D6'):
    tblpr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:color'), color)
        b.append(el)
    tblpr.append(b)


def no_borders(table):
    tblpr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'nil')
        b.append(el)
    tblpr.append(b)


def cell_margins(table, top=80, bottom=80, left=110, right=110):
    tblpr = table._tbl.tblPr
    m = OxmlElement('w:tblCellMar')
    for edge, v in (('top', top), ('bottom', bottom), ('left', left), ('right', right)):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:w'), str(v))
        el.set(qn('w:type'), 'dxa')
        m.append(el)
    tblpr.append(m)


def full_width(table):
    """Pins a table to the text width, centred, so it lines up with the margins."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblpr = table._tbl.tblPr
    w = OxmlElement('w:tblW')
    w.set(qn('w:w'), str(int(CONTENT_W.twips)))
    w.set(qn('w:type'), 'dxa')
    for old in tblpr.findall(qn('w:tblW')):
        tblpr.remove(old)
    tblpr.append(w)
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblpr.append(layout)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------- helpers the content calls
mm = 1
W = 210
MARGIN = 18


def p(text, style='body'):
    par = doc.add_paragraph()
    if style == 'toc':
        par.paragraph_format.space_after = Pt(3)
        rich(par, text, size=11)
    elif style == 'small':
        rich(par, text, size=8.5, color=GREY)
    else:
        rich(par, text)


def h1(text):
    rich(doc.add_paragraph(style='Heading 1'), text)


def h2(text):
    rich(doc.add_paragraph(style='Heading 2'), text)


def h3(text):
    rich(doc.add_paragraph(style='Heading 3'), text)


def _list(items, marker):
    for i, it in enumerate(items, 1):
        par = doc.add_paragraph()
        fmt = par.paragraph_format
        fmt.left_indent = Mm(6.5)
        fmt.first_line_indent = Mm(-5)
        fmt.space_after = Pt(3)
        fmt.tab_stops.add_tab_stop(Mm(6.5))
        par.add_run((marker(i)) + '\t')
        rich(par, it)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def bullets(items):
    _list(items, lambda i: '•')


def steps(items):
    _list(items, lambda i: f'{i}.')


def box(title, lines, tone='tip'):
    fill, title_color = ('E0F1DC', RGBColor(0x2F, 0x5A, 0x2A)) if tone == 'tip' else ('FFF4E0', RGBColor(0x8A, 0x5A, 0x00))
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_borders(t)
    full_width(t)
    cell_margins(t, 120, 120, 160, 160)
    c = t.rows[0].cells[0]
    c.width = CONTENT_W
    shade(c, fill)
    first = c.paragraphs[0]
    run = first.add_run(title)
    run.bold = True
    run.font.color.rgb = title_color
    for line in lines:
        par = c.add_paragraph()
        par.paragraph_format.space_after = Pt(3)
        rich(par, line, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def table(rows, widths, header=True):
    total = sum(widths)
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    full_width(t)
    borders(t)
    cell_margins(t)
    col_w = [int(CONTENT_W * (w / total)) for w in widths]
    for r, row in enumerate(rows):
        for ci, text in enumerate(row):
            cell = t.rows[r].cells[ci]
            cell.width = col_w[ci]
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            if header and r == 0:
                shade(cell, 'FAF6F0')
                rich(par, f'<b>{text}</b>' if text else '', size=9)
            else:
                rich(par, text, size=9)
    if header:
        trpr = t.rows[0]._tr.get_or_add_trPr()
        el = OxmlElement('w:tblHeader')
        el.set(qn('w:val'), 'true')
        trpr.append(el)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def link(url, text=None):
    return f'<link href="{url}" color="#9B2C40"><u>{text or url}</u></link>'


class Story:
    """The content's direct story.append(...) calls: page breaks matter, the rest is spacing."""

    def append(self, item):
        if type(item).__name__ == 'PageBreak':
            page_break()


class PageBreak:
    pass


class NextPageTemplate:
    def __init__(self, *_):
        pass


class Spacer:
    def __init__(self, *_):
        pass


story = Story()

# ---------------------------------------------------------------- cover
band = doc.add_table(rows=1, cols=1)
no_borders(band)
full_width(band)
cell_margins(band, 900, 900, 400, 400)
c = band.rows[0].cells[0]
c.width = CONTENT_W
shade(c, '9B2C40')
for text, size, bold, after in [
    ('BALIYE COUTURE', 12, False, 14),
    ('Admin Panel Guide', 30, True, 6),
    ('AI Studio, Design Builder and everything else', 14, False, 0),
]:
    par = c.paragraphs[0] if text == 'BALIYE COUTURE' else c.add_paragraph()
    par.paragraph_format.space_after = Pt(after)
    run = par.add_run(text)
    run.bold, run.font.size = bold, Pt(size)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

doc.add_paragraph().paragraph_format.space_after = Pt(30)
for line in [
    'For the Baliye Couture team.',
    'How to sign in, create AI embroidery, set up Create Your Own Design, and look after the website content.',
]:
    rich(doc.add_paragraph(), line, size=11)
doc.add_paragraph().paragraph_format.space_after = Pt(10)
for label, url in [('WEBSITE', SITE_URL), ('ADMIN PANEL', ADMIN_URL)]:
    lab = doc.add_paragraph()
    lab.paragraph_format.space_after = Pt(0)
    r = lab.add_run(label)
    r.font.size, r.font.color.rgb = Pt(9), GREY
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(12)
    add_hyperlink(par, url, url, bold=True)
    for run in par._p.iter(qn('w:r')):
        rpr = run.find(qn('w:rPr'))
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), '26')
        rpr.append(sz)
doc.add_paragraph().paragraph_format.space_after = Pt(120)
v = doc.add_paragraph()
r = v.add_run('Version: 30 October 2026')
r.font.size, r.font.color.rgb = Pt(9), GREY

# ---------------------------------------------------------------- header / footer with page numbers
new = doc.add_section(WD_SECTION.NEW_PAGE)
new.header.is_linked_to_previous = False
new.footer.is_linked_to_previous = False
hp = new.header.paragraphs[0]
hr = hp.add_run('Baliye Couture  ·  Admin Panel Guide')
hr.font.size, hr.font.color.rgb = Pt(8), GREY
fp = new.footer.paragraphs[0]
fp.alignment = 2
fr = fp.add_run('Page ')
fr.font.size, fr.font.color.rgb = Pt(8), GREY
for kind, text in (('begin', None), (None, 'PAGE'), ('end', None)):
    run = fp.add_run()
    run.font.size, run.font.color.rgb = Pt(8), GREY
    if kind:
        fc = OxmlElement('w:fldChar')
        fc.set(qn('w:fldCharType'), kind)
        run._r.append(fc)
    else:
        instr = OxmlElement('w:instrText')
        instr.set(qn('xml:space'), 'preserve')
        instr.text = text
        run._r.append(instr)

# ---------------------------------------------------------------- the content, from admin-panel-guide.py
start = SOURCE.index('# ---------------------------------------------------------------- contents')
end = SOURCE.index('# ---------------------------------------------------------------- build')
content = SOURCE[start:end]
# The section already starts on a new page; drop the PDF's opening page break.
content = content.replace("story.append(NextPageTemplate('body'))\nstory.append(PageBreak())\n", '', 1)
exec(compile(content, 'admin-panel-guide.py (content)', 'exec'))

doc.core_properties.title = 'Baliye Couture Admin Panel Guide'
doc.core_properties.author = 'Baliye Couture'
doc.save(OUT)
print('written', OUT)
